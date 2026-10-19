"""Bounded, private previews and atomic imports. No original upload is retained."""
import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone

from apps.billing.services import require_access
from apps.core.models import Membership, Organization
from apps.legal.forms import ContactForm, CourtCaseForm
from apps.legal.models import Contact, CourtCase
from apps.legal.services import audit, visible_matters
from .models import ImportBatch


class PreviewUnavailable(ValidationError):
    """A stale confirmation can never revive a discarded or expired preview."""

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 12
MAX_CELL = 4000
MAX_EXPANDED = 12 * 1024 * 1024
FIELDS = {
    "contacts": [("name", "Nome", ["nome", "nome razao social", "razao social"]), ("kind", "Tipo de pessoa", ["tipo", "tipo de pessoa"]), ("document", "CPF / CNPJ", ["cpf", "cnpj", "cpf cnpj", "documento"]), ("email", "E-mail", ["e mail", "email"]), ("phone", "Telefone", ["telefone", "whatsapp", "celular"]), ("birth_date", "Nascimento", ["nascimento", "data de nascimento"] )],
    "cases": [("matter", "ID da pasta", ["pasta", "id pasta", "id da pasta", "pasta id"]), ("number", "Número do processo", ["numero", "numero do processo", "processo", "numero cnj"]), ("court", "Tribunal / órgão", ["tribunal", "orgao", "tribunal orgao"]), ("opposing_party", "Parte contrária", ["parte contraria"]), ("notes", "Observações", ["observacoes", "anotacoes"])]
}


def normalized(value):
    return " ".join(re.sub(r"[^a-z0-9]+", " ", unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().lower()).split())


def authorize(actor, organization, *, lock=False):
    if not actor or not actor.is_authenticated or not actor.is_active or not Organization.objects.filter(pk=organization.pk, status="active", business_type="office").exists():
        raise PermissionDenied("A importação exige um escritório jurídico ativo.")
    memberships = Membership.objects.filter(organization=organization, user=actor, is_active=True, user__is_active=True, role__in=["owner", "manager"])
    if lock:
        memberships = memberships.select_for_update()
    member = memberships.first()
    if not member:
        raise PermissionDenied("Somente a direção e a gestão podem importar cadastros.")
    from apps.core.services import member_has_module
    if not member_has_module(member, "intake"):
        raise PermissionDenied("Seu acesso à importação de planilhas foi restringido.")
    require_access(organization)


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    value = str(value).strip()
    if len(value) > MAX_CELL or "\x00" in value:
        raise ValidationError("Cada célula deve ter até 4.000 caracteres, sem caracteres nulos.")
    if value.startswith(("=", "@")) or (value.startswith(("+", "-")) and not re.fullmatch(r"[+-]\d[\d\s().-]*", value)):
        raise ValidationError("A planilha contém fórmulas ou comandos. Cole apenas os valores e envie novamente.")
    return value


def _xlsx_rows(raw):
    from defusedxml.ElementTree import iterparse
    from openpyxl import load_workbook
    from openpyxl.xml import DEFUSEDXML
    if not DEFUSEDXML:
        raise ValidationError("O leitor seguro de planilhas não está disponível.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(e.file_size for e in entries) > MAX_EXPANDED:
                raise ValidationError("O XLSX excede o limite de conteúdo expandido. Divida a planilha.")
            nodes = 0
            for entry in entries:
                lower = entry.filename.lower()
                if "externallinks/" in lower or "vbaproject" in lower or entry.flag_bits & 1:
                    raise ValidationError("Não são aceitos arquivos com macros, vínculos externos ou criptografia.")
                if lower.endswith((".xml", ".rels")):
                    for _, element in iterparse(io.BytesIO(archive.read(entry)), events=("end",)):
                        nodes += 1
                        if nodes > 100000:
                            raise ValidationError("O XLSX possui estrutura excessiva. Use o modelo simples disponibilizado.")
                        if element.tag.endswith("}Relationship") and element.attrib.get("TargetMode") == "External":
                            raise ValidationError("Remova os vínculos externos da planilha antes de importar.")
                        element.clear()
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            result = None
            for sheet in workbook.worksheets:
                # Reset dimensions so forged worksheet bounds cannot hide cells.
                sheet.reset_dimensions()
                rows = []
                for line, cells in enumerate(sheet.iter_rows(), start=1):
                    if line > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                        raise ValidationError("Use até 500 linhas de dados e 12 colunas por arquivo, sem formatação fora desse intervalo.")
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValidationError("A planilha contém fórmulas ou erros de cálculo. Cole apenas os valores.")
                    rows.append([cell.value for cell in cells])
                if any(any(cell is not None for cell in row) for row in rows):
                    if result is not None:
                        raise ValidationError("Use somente uma aba preenchida por arquivo.")
                    result = rows
            return result or []
        finally:
            workbook.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("Não foi possível ler o XLSX. Salve uma planilha simples, sem senha, ou utilize o CSV modelo.") from exc


def parse_upload(upload, kind):
    if kind not in FIELDS:
        raise ValidationError("Selecione Pessoas ou Processos.")
    if not upload.size or upload.size > MAX_BYTES:
        raise ValidationError("O arquivo deve ter conteúdo e no máximo 2 MB.")
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError("O arquivo deve ter no máximo 2 MB.")
    extension = Path(upload.name).suffix.lower()
    if extension == ".xlsx":
        source = _xlsx_rows(raw)
    elif extension == ".csv":
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeError as exc:
            raise ValidationError("Salve o CSV com codificação UTF-8 e envie novamente.") from exc
        try:
            delimiter = csv.Sniffer().sniff(decoded[:8192], delimiters=";,\t").delimiter
        except csv.Error:
            delimiter = ";"
        try:
            source = []
            for row in csv.reader(io.StringIO(decoded, newline=""), delimiter=delimiter, strict=True):
                source.append(row)
                if len(source) > MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                    raise ValidationError("Use até 500 linhas de dados e 12 colunas por arquivo.")
        except csv.Error as exc:
            raise ValidationError("CSV inválido. Confira as aspas e os separadores ou use o modelo.") from exc
    else:
        raise ValidationError("Envie um arquivo CSV ou XLSX.")
    if not source:
        raise ValidationError("A planilha está vazia.")
    aliases = {normalized(alias): (field, label) for field, label, names in FIELDS[kind] for alias in [field, label, *names]}
    columns, indexes, seen = [], [], set()
    for index, value in enumerate(source[0]):
        header = _cell(value)
        if len(header) > 80:
            raise ValidationError("Os cabeçalhos devem ter até 80 caracteres. Use os nomes das colunas do modelo.")
        if not header:
            if any(index < len(row) and row[index] not in (None, "") for row in source[1:]):
                raise ValidationError("Toda coluna preenchida precisa de um cabeçalho.")
            continue
        match = aliases.get(normalized(header))
        if not match:
            raise ValidationError(f"Coluna não reconhecida: {header[:80]}. Use os cabeçalhos do modelo.")
        field, label = match
        if field in seen:
            raise ValidationError(f"A coluna {label} aparece mais de uma vez.")
        seen.add(field)
        columns.append({"source": header, "field": field, "label": label})
        indexes.append((index, field))
    required = {"name"} if kind == "contacts" else {"matter", "number"}
    if not required.issubset(seen):
        raise ValidationError("Faltam colunas obrigatórias. Pessoas exige Nome; Processos exige ID da pasta e Número do processo.")
    rows = []
    for line, source_row in enumerate(source[1:], start=2):
        if not any(value not in (None, "") for value in source_row):
            continue
        if len(source_row) > len(source[0]):
            raise ValidationError(f"Linha {line}: existem valores além dos cabeçalhos.")
        data = {}
        for index, field in indexes:
            raw_value = source_row[index] if index < len(source_row) else ""
            if field in {"document", "number", "phone"} and isinstance(raw_value, (int, float)):
                raise ValidationError(f"Linha {line}: formate CPF, CNPJ, telefone e número do processo como texto para preservar todos os dígitos.")
            data[field] = _cell(raw_value)
        if kind == "contacts":
            kind_value = normalized(data.get("kind", ""))
            data["kind"] = {"": "person", "person": "person", "pessoa fisica": "person", "pf": "person", "company": "company", "pessoa juridica": "company", "pj": "company"}.get(kind_value, data.get("kind", ""))
            if data.get("birth_date"):
                try:
                    data["birth_date"] = datetime.strptime(data["birth_date"], "%d/%m/%Y").date().isoformat()
                except ValueError:
                    pass  # The normal ContactForm reports invalid dates.
        rows.append({"line": line, "data": data})
    if not rows:
        raise ValidationError("Inclua pelo menos uma linha de dados abaixo dos cabeçalhos.")
    return rows, columns


def _contact_keys(data):
    # Keep non-Latin names distinct while folding case, accents and spacing.
    name = "".join(char for char in unicodedata.normalize("NFKD", data.get("name", "")) if not unicodedata.combining(char)).casefold()
    name = " ".join(re.sub(r"[^\w]+", " ", name, flags=re.UNICODE).split())
    keys = [("name", name, data.get("kind") or "person")]
    if data.get("document"):
        keys.append(("document", re.sub(r"\W+", "", data["document"]).casefold()))
    if data.get("email"):
        keys.append(("email", data["email"].strip().casefold()))
    return keys


def validate_rows(*, actor, organization, kind, rows):
    authorize(actor, organization)
    if kind == "contacts":
        existing = set()
        for contact in Contact.objects.filter(organization=organization).values("name", "kind", "document", "email").iterator():
            existing.update(_contact_keys(contact))
        form_class = ContactForm
    else:
        existing = {normalized(number).replace(" ", "") for number in CourtCase.objects.filter(organization=organization).values_list("number", flat=True).iterator()}
        visible_existing = {normalized(number).replace(" ", "") for number in CourtCase.objects.filter(organization=organization, matter__in=visible_matters(actor, organization)).values_list("number", flat=True).iterator()}
        form_class = CourtCaseForm
    seen, result, forms = set(), [], []
    for row in rows:
        data = row["data"]
        form = form_class(data, organization=organization, user=actor)
        valid = form.is_valid()
        keys = _contact_keys(data) if kind == "contacts" else [normalized(data.get("number", "")).replace(" ", "")]
        restricted_conflict = kind == "cases" and keys[0] in existing and keys[0] not in visible_existing
        errors = []
        for field, error_list in form.errors.as_data().items():
            for error in error_list:
                # Global number uniqueness must survive confidentiality changes,
                # but its standard model error must not expose hidden records.
                if restricted_conflict and error.code in {"unique", "unique_together"}:
                    continue
                errors.extend(f"{form.fields[field].label if field in form.fields else 'Registro'}: {message}" for message in error.messages)
        if any(key in existing for key in keys):
            errors.append("Este número não está disponível para importação. Confira o preenchimento com a administração do escritório." if restricted_conflict else "Possível duplicidade no escritório. Confira o cadastro existente; nenhuma informação será sobrescrita.")
        if any(key in seen for key in keys):
            errors.append("Possível duplicidade dentro desta planilha.")
        seen.update(keys)
        values = []
        for field, label, _ in FIELDS[kind]:
            value = data.get(field, "") or "—"
            if field == "kind":
                value = dict(Contact.Kind.choices).get(value, value)
            if field == "birth_date" and form.cleaned_data.get(field):
                value = form.cleaned_data[field].strftime("%d/%m/%Y")
            if field == "matter" and form.cleaned_data.get(field):
                value = f"#{form.cleaned_data[field].pk} · {form.cleaned_data[field].title}"
            values.append({"label": label, "value": value})
        item = {"line": row["line"], "title": data.get("name") or data.get("number") or "Sem identificação", "values": values, "errors": errors, "valid": valid and not errors}
        result.append(item)
        forms.append(form)
    invalid = sum(not item["valid"] for item in result)
    return {"rows": result, "valid_count": len(result) - invalid, "invalid_count": invalid, "can_confirm": bool(result) and invalid == 0}, forms


@transaction.atomic
def stage_import(*, actor, organization, kind, upload):
    Organization.objects.select_for_update().get(pk=organization.pk)
    authorize(actor, organization, lock=True)
    # Bound retained previews per account and erase payloads once expired.
    ImportBatch.objects.filter(organization=organization, created_by=actor, expires_at__lte=timezone.now(), committed_at__isnull=True).delete()
    if ImportBatch.objects.filter(organization=organization, created_by=actor, committed_at__isnull=True).count() >= 10:
        raise ValidationError("Você já tem 10 prévias abertas. Descarte uma prévia antes de enviar outro arquivo.")
    rows, columns = parse_upload(upload, kind)
    return ImportBatch.objects.create(organization=organization, created_by=actor, kind=kind, rows=rows, columns=columns, row_count=len(rows), expires_at=timezone.now() + timedelta(hours=24))


@transaction.atomic
def commit_import(*, actor, organization, batch_id):
    # Every importer for the office shares this lock, including repeated clicks.
    Organization.objects.select_for_update().get(pk=organization.pk)
    authorize(actor, organization, lock=True)
    try:
        batch = ImportBatch.objects.select_for_update().get(pk=batch_id, organization=organization, created_by=actor)
    except ImportBatch.DoesNotExist as exc:
        raise PreviewUnavailable("Esta prévia não está mais disponível. Envie novamente sua planilha.") from exc
    if batch.committed_at:
        return batch, False
    if batch.expires_at <= timezone.now():
        raise PreviewUnavailable("Esta prévia expirou. Envie novamente a planilha para conferir os dados atuais.")
    review, forms = validate_rows(actor=actor, organization=organization, kind=batch.kind, rows=batch.rows)
    if not review["can_confirm"]:
        raise ValidationError("Os dados ou suas permissões mudaram. Confira os problemas da prévia; nenhuma linha foi importada.")
    for form in forms:
        try:
            obj = form.save()
        except (IntegrityError, ValidationError) as exc:
            # A simultaneous ordinary case registration also uses the existing
            # database constraint; roll the entire batch back on collision.
            raise ValidationError("Não foi possível criar o lote com os dados atuais. Nenhuma linha foi importada; confira a prévia e tente novamente.") from exc
        audit(organization=organization, actor=actor, action="record.imported", obj=obj, description="Cadastro criado por importação assistida conferida.")
        if batch.kind == "contacts" and isinstance(obj, Contact):
            from apps.legal.services import ensure_client_matter
            ensure_client_matter(contact=obj, actor=actor)
    batch.created_count = len(forms)
    batch.committed_at = timezone.now()
    batch.rows = []  # Imported data now exists in the regular permission-controlled records.
    batch.save(update_fields=["created_count", "committed_at", "rows"])
    audit(organization=organization, actor=actor, action="workspace.imported", description=f"Importação assistida: {batch.created_count} registros de {batch.get_kind_display().lower()}.")
    return batch, True
