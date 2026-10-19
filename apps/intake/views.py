import csv
import io

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db.models import Q
from django.db import transaction
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_http_methods, require_POST

from apps.legal.services import visible_matters
from apps.core.models import Organization
from apps.legal.views import organization_required
from .forms import ConfirmationForm, UploadForm
from .models import ImportBatch
from .services import FIELDS, PreviewUnavailable, authorize, commit_import, stage_import, validate_rows


def _batch(request, pk):
    authorize(request.user, request.organization)
    return get_object_or_404(ImportBatch, pk=pk, organization=request.organization, created_by=request.user)


@never_cache
@organization_required
@require_http_methods(["GET", "POST"])
def home(request):
    authorize(request.user, request.organization)
    form = UploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            batch = stage_import(actor=request.user, organization=request.organization, kind=form.cleaned_data["kind"], upload=form.cleaned_data["file"])
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            return redirect("intake:preview", pk=batch.pk)
    query = request.GET.get("matter_q", "").strip()[:100]
    matters = visible_matters(request.user, request.organization).select_related("client").order_by("title")
    if query:
        matters = matters.filter(Q(title__icontains=query) | Q(client__name__icontains=query) | Q(pk=int(query) if query.isdecimal() and len(query) < 12 else -1))
    return render(request, "intake/home.html", {"page_title": "Importação assistida", "active_nav": "intake", "form": form, "batches": ImportBatch.objects.filter(organization=request.organization, created_by=request.user).filter(Q(committed_at__isnull=False) | Q(expires_at__gt=timezone.now()))[:10], "matters": matters[:20], "matter_q": query})


def _render_preview(request, batch, *, confirmation=None, status=200):
    expired = not batch.committed_at and batch.expires_at <= timezone.now()
    review = {"rows": [], "can_confirm": False, "valid_count": 0, "invalid_count": 0}
    if not expired and not batch.committed_at:
        review, _ = validate_rows(actor=request.user, organization=request.organization, kind=batch.kind, rows=batch.rows)
    page = Paginator(review.pop("rows"), 25).get_page(request.GET.get("page"))
    return render(request, "intake/preview.html", {"page_title": "Conferir importação", "active_nav": "intake", "batch": batch, "expired": expired, "review": review, "page": page, "confirmation": confirmation or ConfirmationForm()}, status=status)


@never_cache
@organization_required
@require_GET
def preview(request, pk):
    return _render_preview(request, _batch(request, pk))


@never_cache
@organization_required
@require_POST
def confirm(request, pk):
    batch = _batch(request, pk)
    form = ConfirmationForm(request.POST)
    if not form.is_valid():
        return _render_preview(request, batch, confirmation=form, status=400)
    try:
        batch, created = commit_import(actor=request.user, organization=request.organization, batch_id=batch.pk)
    except PreviewUnavailable as exc:
        messages.error(request, " ".join(exc.messages))
        return redirect("intake:home")
    except ValidationError as exc:
        form.add_error(None, exc)
        return _render_preview(request, batch, confirmation=form, status=400)
    suffix = "s" if batch.created_count != 1 else ""
    if batch.kind == "contacts":
        messages.success(request, f"{batch.created_count} pessoa{suffix} importada{suffix} com sucesso e suas pastas jurídicas vinculadas criadas automaticamente." if created else "Este lote já foi importado. Nenhum registro foi duplicado.")
    else:
        messages.success(request, f"{batch.created_count} registro{suffix} importado{suffix} com sucesso." if created else "Este lote já foi importado. Nenhum registro foi duplicado.")
    return redirect("intake:preview", pk=batch.pk)


@never_cache
@organization_required
@require_POST
@transaction.atomic
def discard(request, pk):
    Organization.objects.select_for_update().get(pk=request.organization.pk)
    batch = _batch(request, pk)
    if batch.committed_at:
        messages.info(request, "O comprovante de uma importação concluída é preservado.")
    else:
        batch.delete()
        messages.success(request, "Prévia descartada. Os cadastros do escritório não foram alterados.")
    return redirect("intake:home")


@never_cache
@organization_required
@require_GET
def template(request, kind, extension):
    authorize(request.user, request.organization)
    if kind not in FIELDS or extension not in {"csv", "xlsx"}:
        raise Http404
    headers = [label for _, label, _ in FIELDS[kind]]
    if extension == "csv":
        output = io.StringIO(newline="")
        csv.writer(output, delimiter=";", lineterminator="\r\n").writerow(headers)
        response = HttpResponse(output.getvalue().encode("utf-8-sig"), content_type="text/csv; charset=utf-8")
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pessoas" if kind == "contacts" else "Processos"
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(1, index)
            cell.font = Font(bold=True, color="171A20")
            cell.fill = PatternFill("solid", fgColor="F5A000")
            sheet.column_dimensions[get_column_letter(index)].width = max(23, len(header) + 5)
            sheet.cell(2, index).number_format = "@"
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="jumbi-modelo-{kind}.{extension}"'
    response["X-Content-Type-Options"] = "nosniff"
    return response
